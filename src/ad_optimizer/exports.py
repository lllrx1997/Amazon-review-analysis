from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import Finding, Report


def count(item) -> int:
    return len({e.review_id for e in item.evidence})


def topic_counts(topic, report: Report) -> list[int]:
    indexed = {r.id: r.rating for r in report.reviews}
    ids = {e.review_id for e in topic.evidence}
    return [sum(indexed[i] in group for i in ids) for group in ({1, 2}, {3}, {4, 5})]


def workbook_bytes(report: Report) -> bytes:
    book = Workbook()
    book.remove(book.active)

    def sheet(name, headers, rows):
        page = book.create_sheet(name)
        page.append(headers)
        for row in rows:
            # Excel has a 32767-character cell limit. Split rather than truncate.
            values = [str(v) if v is not None else "" for v in row]
            chunks = max([1] + [(len(v) + 29999) // 30000 for v in values])
            for part in range(chunks):
                page.append([v[part * 30000:(part + 1) * 30000] if len(v) > 30000 else (v if part == 0 or i == 0 else "") for i, v in enumerate(values)])
        for row in page:
            for cell in row:
                if isinstance(cell.value, str):
                    # Prevent formula execution, including strings beginning '='.
                    cell.value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", cell.value)
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in page[1]:
            cell.fill = PatternFill("solid", fgColor="172B4D")
            cell.font = Font(color="FFFFFF", bold=True)
        page.freeze_panes = "A2"
        page.auto_filter.ref = page.dimensions
        for col in page.columns:
            page.column_dimensions[col[0].column_letter].width = 28 if len(headers) > 3 else 40

    summary = [
        ("ASIN", report.asin), ("商品", report.product_name or "数据源未提供"), ("商品链接", report.product_url),
        ("来源", report.source), ("评论范围", report.scope), ("采集时间 UTC", report.collected_at.isoformat()),
        ("生成时间 UTC", report.generated_at.isoformat()), ("模型", report.model), ("分析规则", report.rule_version),
        ("有效样本", report.total), ("已分析", report.analyzed_count),
        ("样本差评占比", f"{report.negative_percentage}%" if report.negative_percentage is not None else "—"),
        *[(f"{i}星", report.stars[str(i)]) for i in range(1, 6)],
        ("原始返回", report.received_count), ("排除", report.excluded_count), ("重复", report.duplicate_count),
        ("超出分析上限", report.capped_count), ("样本范围说明", "差评占比仅代表当前样本，不代表商品整体差评率"),
        *[("限制说明", warning) for warning in report.warnings],
        *[("好评摘要", f"{f.name}：{f.description}；证据：{','.join(dict.fromkeys(e.review_id for e in f.evidence))}") for f in report.praises],
    ]
    sheet("概览", ["项目", "内容"], summary)

    def finding_rows(items):
        return [[f.name, count(f), f.description, f"AI推测，尚需验证：{f.possible_cause}" if f.possible_cause else "", "\n".join(f"{e.review_id}: {e.quote}" for e in f.evidence)] for f in items]

    sheet("差评问题", ["问题", "涉及条数", "买家明确反馈", "可能原因（非确定结论）", "原文证据"], finding_rows(report.issues))
    sheet("用户认可", ["认可点", "涉及条数", "反馈", "备注", "原文证据"], finding_rows(report.praises))
    sheet("三星反馈", ["主题", "涉及条数", "认可、不足或诉求", "备注", "原文证据"], finding_rows(report.neutral_feedback))
    sheet("评价主题", ["主题", "总条数", "1–2星", "3星", "4–5星", "评论ID"], [[t.name, count(t), *topic_counts(t, report), ",".join(dict.fromkeys(e.review_id for e in t.evidence))] for t in report.topics])
    sheet("评论明细", ["ID", "请求ASIN", "来源ASIN", "星级", "日期", "标题原文", "正文原文", "标题中文", "正文中文", "评论链接", "分析状态", "主题"], [[r.id, r.requested_asin, r.source_asin or "未确认", r.rating, r.date or r.date_raw, r.title_original, r.body_original, r.title_zh, r.body_zh, r.review_url, "完成" if r.analyzed else "未完成", ",".join(t.name for t in report.topics if r.id in {e.review_id for e in t.evidence})] for r in report.reviews])
    if report.usage_insights:
        sheet("使用情境", ["情境", "涉及条数", "原文支持的洞察", "备注", "证据"], finding_rows(report.usage_insights))
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def md(value: str) -> str:
    return re.sub(r"([\\`*_{}\[\]<>()#!|~])", r"\\\1", value).replace("\r", "")


def markdown_bytes(report: Report) -> bytes:
    lines = [f"# 亚马逊评论分析：{report.asin}", "", f"商品：{md(report.product_name or '数据源未提供名称')}", "", f"商品地址：{report.product_url}", "", f"来源：{report.source}；范围：{report.scope}；模型：{md(report.model)}；规则：{report.rule_version}", "", f"采集时间（UTC）：{report.collected_at.isoformat()}；生成时间（UTC）：{report.generated_at.isoformat()}", "", "## 样本概览", "", f"有效样本 {report.total} 条；已分析 {report.analyzed_count} 条。", "", "；".join(f"{i}星 {report.stars[str(i)]}条" for i in range(1, 6)), "", f"样本差评占比：{str(report.negative_percentage) + '%' if report.negative_percentage is not None else '—'}。仅代表当前样本，不代表商品整体差评率。", "", "## 数据限制", "", *[f"- {md(w)}" for w in report.warnings]]
    for title, items in (("差评问题与可能原因", report.issues), ("用户认可点与好评摘要（4–5星）", report.praises), ("3星反馈", report.neutral_feedback), ("使用情境洞察", report.usage_insights)):
        lines += ["", f"## {title}", ""]
        if not items:
            lines.append("当前样本无可展示的已核验反馈；若分析未完成，不代表没有相关问题。")
        for finding in items:
            lines += [f"### {md(finding.name)}（{count(finding)}条）", "", md(finding.description), ""]
            if finding.possible_cause:
                lines += [f"**AI推测，尚需验证**：{md(finding.possible_cause)}", ""]
            for evidence in finding.evidence:
                lines += [f"- 证据 {evidence.review_id}：{md(evidence.quote)}"]
    lines += ["", "## 评价主题", "", "|主题|总条数|1–2星|3星|4–5星|", "|---|---:|---:|---:|---:|"]
    lines += [f"|{md(t.name)}|{count(t)}|{'|'.join(map(str, topic_counts(t, report)))}|" for t in report.topics]
    lines += ["", "同一评论可以涉及多个主题。", "", "## 全部评论证据附录", ""]
    for review in report.reviews:
        lines += [f"### {review.id}", "", f"请求ASIN：{review.requested_asin}；来源ASIN：{review.source_asin or '未确认'}；{review.rating}星；日期：{md(review.date or review.date_raw or '未知')}", "", f"原评论链接：{md(review.review_url or '未提供')}", "", "**原文**", "", md(review.title_original), "", md(review.body_original), "", "**中文翻译**", "", md(review.title_zh), "", md(review.body_zh or '未完成翻译'), ""]
    return "\n".join(lines).encode("utf-8")
