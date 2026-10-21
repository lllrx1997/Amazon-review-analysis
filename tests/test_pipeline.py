import asyncio
from datetime import timedelta
from io import BytesIO
from time import sleep

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from ad_optimizer.analysis import Analyzer, merge_findings, validate_extraction
from ad_optimizer.api import create_app
from ad_optimizer.config import Settings
from ad_optimizer.exports import markdown_bytes, topic_counts, workbook_bytes
from ad_optimizer.models import Evidence, Extraction, Finding, Review, Task, Topic, Translation, utcnow
from ad_optimizer.providers import BrightData, ServiceError, normalize

ASIN = "B012345678"


@pytest.fixture
def settings():
    return Settings(_env_file=None, bright_data_api_key="test-only", bright_data_dataset_id="test-dataset", llm_api_key="test-only", llm_model="test-model")


def sample_rows():
    return [{"review_id": f"id-{i}", "asin": ASIN, "rating": i, "review_text": f"Review number {i}. It works.", "review_date": f"October {i}, 2025", "product_name": "Test product"} for i in range(1, 6)]


def extraction_for(reviews):
    def finding(review):
        return Finding(name="功能体验", description="买家描述使用体验", evidence=[Evidence(review_id=review.id, quote=review.body_original)])
    return Extraction(
        translations=[Translation(review_id=r.id, title_zh="", body_zh=f"第{r.rating}条评论。可以工作。") for r in reviews],
        issues=[finding(r) for r in reviews if r.rating <= 2],
        praises=[finding(r) for r in reviews if r.rating >= 4],
        neutral_feedback=[finding(r) for r in reviews if r.rating == 3],
        topics=[Topic(name="功能体验", evidence=[Evidence(review_id=r.id, quote=r.body_original) for r in reviews])],
    )


def test_normalization_counts_dedup_and_order(settings):
    rows = sample_rows()
    rows += [rows[0], {"rating": "4.5", "review_text": "invalid"}, {"rating": 5, "review_text": ""}]
    report = normalize(rows, ASIN, settings)
    assert report.total == 5
    assert report.duplicate_count == 1
    assert report.excluded_count == 2
    assert report.negative_percentage == 40.0
    assert report.stars == {str(i): 1 for i in range(1, 6)}
    assert report.reviews[0].rating == 5
    assert not report.latest_verified and report.warnings


def test_source_not_assumed_and_cap_warning(settings):
    settings.max_reviews = 1
    rows = sample_rows()
    for row in rows:
        row.pop("asin")
    report = normalize(rows, ASIN, settings)
    assert report.total == 1 and report.capped_count == 4
    assert report.reviews[0].source_asin is None
    assert not report.limit_verified


def test_empty_is_not_zero_percent(settings):
    report = normalize([], ASIN, settings)
    assert report.total == 0 and report.negative_percentage is None


@pytest.mark.parametrize("date_value", ["October 5", "not-a-date", "2099-01-01"])
def test_missing_or_invalid_dates_not_fabricated(settings, date_value):
    rows = sample_rows()[:1]
    rows[0]["review_date"] = date_value
    report = normalize(rows, ASIN, settings)
    assert report.reviews[0].date is None


def test_evidence_rejects_wrong_star_and_fake_quote(settings):
    report = normalize(sample_rows(), ASIN, settings)
    extraction = extraction_for(report.reviews)
    validate_extraction(extraction, report.reviews)
    extraction.issues[0].evidence[0].quote = "This was never said."
    with pytest.raises(ValueError):
        validate_extraction(extraction, report.reviews)
    extraction = extraction_for(report.reviews)
    extraction.issues[0].evidence[0] = Evidence(review_id=report.reviews[0].id, quote=report.reviews[0].body_original)
    with pytest.raises(ValueError):
        validate_extraction(extraction, report.reviews)


def test_translation_must_cover_all_ids(settings):
    reviews = normalize(sample_rows(), ASIN, settings).reviews
    extraction = extraction_for(reviews)
    extraction.translations.pop()
    with pytest.raises(ValueError):
        validate_extraction(extraction, reviews)


def test_merged_evidence_counts_distinct_reviews(settings):
    report = normalize(sample_rows(), ASIN, settings)
    extraction = extraction_for(report.reviews)
    merged = merge_findings(extraction.issues + extraction.issues)
    assert len(merged) == 1
    assert len(merged[0].evidence) == 2
    assert topic_counts(extraction.topics[0], report) == [2, 1, 2]


@pytest.mark.asyncio
async def test_same_named_categories_merge_without_another_model_call(settings, monkeypatch):
    report = normalize(sample_rows(), ASIN, settings)
    report.issues = extraction_for(report.reviews).issues

    async def must_not_call(*args):
        raise AssertionError("Same names need no model call")

    monkeypatch.setattr(Analyzer, "json_call", must_not_call)
    await Analyzer(settings).unify(report)
    assert len(report.issues) == 1
    assert len({e.review_id for e in report.issues[0].evidence}) == 2


@pytest.mark.asyncio
async def test_partial_analysis_keeps_full_star_denominator(settings, monkeypatch):
    settings.llm_batch_characters = 1000
    report = normalize(sample_rows(), ASIN, settings)
    # Force each review into a separate batch, without changing evidence semantics.
    for r in report.reviews:
        r.body_original += "x" * 1200

    async def extract(self, batch):
        if batch[0].rating == 1:
            raise ServiceError("LLM_OUTPUT_INVALID", "test failure")
        return extraction_for(batch)

    async def unify(self, report):
        pass

    monkeypatch.setattr(Analyzer, "extract", extract)
    monkeypatch.setattr(Analyzer, "unify", unify)
    await Analyzer(settings).analyze(report)
    assert report.analyzed_count == 4
    assert report.total == 5 and report.negative_percentage == 40.0
    assert any("1 条评论未完成" in w for w in report.warnings)


def test_export_preserves_formula_text_and_long_original(settings):
    report = normalize(sample_rows(), ASIN, settings)
    report.reviews[0].body_original = "=HYPERLINK(\"unsafe\")" + "a" * 34000
    report.reviews[0].body_zh = "翻译"
    book = load_workbook(BytesIO(workbook_bytes(report)))
    sheet = book["评论明细"]
    body = sheet.cell(2, 7)
    assert body.data_type == "s" and body.value.startswith("=HYPERLINK")
    reconstructed = body.value + sheet.cell(3, 7).value
    assert reconstructed == report.reviews[0].body_original
    exported = markdown_bytes(report).decode()
    assert "全部评论证据附录" in exported
    assert "不代表商品整体差评率" in exported
    assert all(r.id in exported for r in report.reviews)


def test_api_missing_configuration_and_validation():
    app = create_app(Settings(_env_file=None))
    with TestClient(app) as client:
        assert client.get("/api/health").json()["configured"] is False
        assert client.post("/api/analyses", json={"asin": "bad"}).status_code == 422
        response = client.post("/api/analyses", json={"asin": ASIN})
        assert response.status_code == 503
        assert "API" in response.json()["detail"]["message"]


def test_full_api_pipeline_and_downloads(settings, monkeypatch):
    async def collect(self, asin):
        return normalize(sample_rows(), asin, settings)

    async def extract(self, batch):
        return extraction_for(batch)

    async def unify(self, report):
        for field in ("issues", "praises", "neutral_feedback", "topics"):
            setattr(report, field, merge_findings(getattr(report, field)))

    monkeypatch.setattr(BrightData, "collect", collect)
    monkeypatch.setattr(Analyzer, "extract", extract)
    monkeypatch.setattr(Analyzer, "unify", unify)
    app = create_app(settings)
    with TestClient(app) as client:
        response = client.post("/api/analyses", json={"asin": f" {ASIN.lower()} "})
        assert response.status_code == 202
        task_id = response.json()["task_id"]
        for _ in range(50):
            task = client.get(f"/api/analyses/{task_id}").json()
            if task["status"] not in ("fetching", "analyzing"):
                break
            sleep(.01)
        assert task["status"] == "completed"
        assert task["result"]["analyzed_count"] == 5
        assert task["revision"] >= 2
        assert "collection_seconds" in task["timings"]
        assert "analysis_seconds" in task["timings"]
        assert "total_seconds" in task["timings"]
        lightweight = client.get(f"/api/analyses/{task_id}?include_result=false").json()
        assert "result" not in lightweight
        assert lightweight["revision"] == task["revision"]
        assert len(task["result"]["neutral_feedback"]) == 1
        assert client.get(f"/api/analyses/{task_id}/export?format=xlsx").content.startswith(b"PK")
        assert "全部评论证据附录" in client.get(f"/api/analyses/{task_id}/export?format=md").text
        assert client.get(f"/api/analyses/{task_id}/export?format=exe").status_code == 422
        app.state.store.tasks[task_id].expires_at = utcnow() - timedelta(seconds=1)
        assert client.get(f"/api/analyses/{task_id}").status_code == 404


def test_busy_rejects_instead_of_queue(settings, monkeypatch):
    async def collect(self, asin):
        await asyncio.sleep(20)
    monkeypatch.setattr(BrightData, "collect", collect)
    with TestClient(create_app(settings)) as client:
        first = client.post("/api/analyses", json={"asin": ASIN})
        assert first.status_code == 202
        second = client.post("/api/analyses", json={"asin": ASIN})
        assert second.status_code == 429


def test_verified_translations_with_dropped_findings_are_partial(settings, monkeypatch):
    async def collect(self, asin):
        return normalize(sample_rows(), asin, settings)

    async def analyze(self, report):
        report.analyzed_count = report.total
        report.analysis_incomplete = True
        report.warnings.append("一项无效反馈已排除，分析不完整。")
        return report

    monkeypatch.setattr(BrightData, "collect", collect)
    monkeypatch.setattr(Analyzer, "analyze", analyze)
    with TestClient(create_app(settings)) as client:
        task_id = client.post("/api/analyses", json={"asin": ASIN}).json()["task_id"]
        for _ in range(50):
            task = client.get(f"/api/analyses/{task_id}").json()
            if task["status"] not in ("fetching", "analyzing"):
                break
            sleep(.01)
        assert task["status"] == "partial"
        assert task["result"]["analysis_incomplete"] is True
        exported = client.get(f"/api/analyses/{task_id}/export?format=md")
        assert "partial" in exported.headers["content-disposition"]


def test_provider_error_is_not_empty_review_result(settings, monkeypatch):
    async def collect(self, asin):
        raise ServiceError("PROVIDER_AUTH_ERROR", "鉴权失败")
    monkeypatch.setattr(BrightData, "collect", collect)
    with TestClient(create_app(settings)) as client:
        task_id = client.post("/api/analyses", json={"asin": ASIN}).json()["task_id"]
        for _ in range(50):
            task = client.get(f"/api/analyses/{task_id}").json()
            if task["status"] == "failed":
                break
            sleep(.01)
        assert task["error"]["code"] == "PROVIDER_AUTH_ERROR"
        assert task["result"] is None
